#!/usr/bin/env python3
"""COMPANYDOCS: the Construction Safety Program (DRAFT) and the 12 toolbox talks as Word files, from the same text the CRM uses
(tools/companydocs_safety.py).

    python3 tools/build_safety_docs.py <out dir> [--address "4014 Broadway Rd, Bartlett, TN 38135"]

Writes "GrayCo Construction Safety Program (DRAFT).docx" and "GrayCo Toolbox Talks (12).docx". The PDFs come from the CRM itself
(node tools/tests/test_COMPANYDOCS.mjs <build> --out <dir>), so they carry exactly what the CRM prints.
Owner-set blanks print as highlighted "[OWNER TO FILL: …]".
"""
import os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import companydocs_safety as S

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = sys.argv[1] if len(sys.argv) > 1 else '.'
ADDR = sys.argv[sys.argv.index('--address') + 1] if '--address' in sys.argv else '4014 Broadway Rd, Bartlett, TN 38135'
LABEL = dict(S.FIELDS)
NAVY = RGBColor(0x0F, 0x24, 0x40)
RED = RGBColor(0xB9, 0x1C, 0x1C)


def prep(t):
    t = t.replace('{company_address}', ADDR)
    return re.sub(r'%%sf:([a-z_]+)%%', lambda m: '[OWNER TO FILL: ' + LABEL.get(m.group(1), m.group(1)) + ']', t)


def runs(p, text, bold=False, size=None, color=None):
    """**bold** and [OWNER TO …] highlighted"""
    text = prep(text)
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        for j, seg in enumerate(re.split(r'(\[OWNER TO [^\]]*\])', part)):
            if not seg:
                continue
            r = p.add_run(seg)
            r.bold = bold or i % 2 == 1
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
            if j % 2 == 1:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def base_doc(footer_text):
    d = Document()
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.7); s.left_margin = s.right_margin = Inches(0.8)
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo = os.path.join(ROOT, 'grayco-logo.jpg')
        if os.path.exists(logo):
            hp.add_run().add_picture(logo, height=Inches(0.55))
        h2 = s.header.add_paragraph(); h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(h2, '**GrayCo Construction Services, LLC**', size=11, color=NAVY)
        h3 = s.header.add_paragraph(); h3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(h3, ADDR + ' · ' + S.PHONE + ' · grayco-cs.com · TN #85034 BCB · AR #0486990627 · MS #27769', size=8)
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(fp, footer_text, size=7.5, color=RED)
    return d


def table_rows(d, rows, widths=(2.4, 4.4)):
    t = d.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for a, b in rows:
        c = t.add_row().cells
        c[0].width = Inches(widths[0]); c[1].width = Inches(widths[1])
        runs(c[0].paragraphs[0], a, size=9.5, color=RGBColor(0x55, 0x55, 0x55))
        runs(c[1].paragraphs[0], '**' + b + '**' if b and '**' not in b else (b or ' '), size=9.5)
    d.add_paragraph()


def sigs(d, a, b):
    t = d.add_table(rows=3, cols=2)
    for i, lab in enumerate(['signature', 'Printed name / title', 'Date']):
        for j, who in enumerate([a, b]):
            p = t.rows[i].cells[j].paragraphs[0]
            p.paragraph_format.space_before = Pt(18)
            runs(p, '_______________________________________\n' + ((who + ' — ' + lab) if lab == 'signature' else lab), size=8.5)
    d.add_paragraph()


def blocks(d, bl):
    for b in bl:
        k = b[0]
        if k == 'draft':
            t = d.add_table(rows=1, cols=1); t.style = 'Table Grid'; c = t.rows[0].cells[0]; shade(c, 'FEF2F2')
            runs(c.paragraphs[0], b[1], size=9.5, color=RGBColor(0x7F, 0x1D, 0x1D)); d.add_paragraph()
        elif k == 'h':
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
            runs(p, b[1], bold=True, size=12.5, color=NAVY)
        elif k == 'h4':
            p = d.add_paragraph(); p.paragraph_format.keep_with_next = True; runs(p, b[1], bold=True, size=10.5, color=RGBColor(0x1E, 0x3A, 0x5F))
        elif k == 'p':
            runs(d.add_paragraph(), b[1])
        elif k == 'note':
            p = d.add_paragraph(); r = runs(p, b[1], size=9.5); [setattr(x, 'italic', True) for x in p.runs]
        elif k == 'ul':
            for x in b[1]:
                runs(d.add_paragraph(style='List Bullet'), x)
        elif k == 'info':
            table_rows(d, b[1])
        elif k == 'lines':
            for x in b[1]:
                if x:
                    runs(d.add_paragraph(), '**' + x + '**')
                else:
                    d.add_paragraph('_' * 95)
        elif k == 'sigs':
            sigs(d, b[1], b[2])
        elif k == 'pb':
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(k)


os.makedirs(out, exist_ok=True)
d = base_doc(S.DRAFT)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, 'Construction Safety Program', bold=True, size=17, color=NAVY)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, S.REV, bold=True, size=9.5, color=RED)
blocks(d, S.SAFETY)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; runs(p, S.REV, size=8.5, color=RGBColor(0x77, 0x77, 0x77))
f1 = os.path.join(out, 'GrayCo Construction Safety Program (DRAFT).docx'); d.save(f1)

d = base_doc('GrayCo toolbox talks — ' + S.TALK_REV)
for i, t in enumerate(S.TALKS):
    if i:
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, 'Toolbox Talk %d — %s' % (i + 1, t[1]), bold=True, size=14, color=NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, '10–15 minutes · read it, then talk about this job', size=9)
    runs(d.add_paragraph(), '**Why it matters:** ' + t[2], size=10)
    runs(d.add_paragraph(), 'Key points', bold=True, size=11, color=NAVY)
    for x in t[3]:
        runs(d.add_paragraph(style='List Bullet'), x, size=10)
    runs(d.add_paragraph(), 'Talk about it on this job', bold=True, size=11, color=NAVY)
    for x in t[4]:
        runs(d.add_paragraph(style='List Bullet'), x, size=10)
    runs(d.add_paragraph(), 'Date: ____________   Job: ______________________________   Led by: ____________________', size=9)
    tb = d.add_table(rows=11, cols=3); tb.style = 'Table Grid'
    for j, h in enumerate(['Printed name', 'Signature', 'Company (GrayCo / sub)']):
        runs(tb.rows[0].cells[j].paragraphs[0], '**' + h + '**', size=9); shade(tb.rows[0].cells[j], 'F1F5F9')
f2 = os.path.join(out, 'GrayCo Toolbox Talks (12).docx'); d.save(f2)
print(f1); print(f2)
